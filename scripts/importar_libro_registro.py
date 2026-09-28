"""Importa una única vez el libro de registro de alumnos desde Excel."""
import argparse
import os
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'autoescuela.settings')

import django

django.setup()

from django.db import transaction

from facturacion.models import Alumno, Autoescuela


def normalizar_dni(value):
    return str(value or '').upper().replace(' ', '').replace('-', '').strip()


def texto(value):
    return str(value or '').strip()


def fecha(value, avisos=None, fila=None, columna=None):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for formato in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), formato).date()
        except ValueError:
            pass
    if avisos is not None:
        avisos.append(f'Fila {fila}, {columna}: fecha no válida "{value}"')
    return None


def nombre_completo(apellido1, apellido2, nombre):
    return ' '.join(parte for parte in (texto(apellido1), texto(apellido2), texto(nombre)) if parte)


def numero_registro(value):
    if value in (None, ''):
        return None
    return int(value)


def importar(path, nombre_autoescuela, dry_run=False):
    autoescuela = Autoescuela.objects.get(nombre=nombre_autoescuela)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        headers = {
            texto(cell.value).upper(): cell.column
            for cell in worksheet[11]
            if texto(cell.value)
        }
        required = {'F_ALTA', 'APELLIDO1', 'NOMBRE', 'DNI/NIF/NIE', 'PMS'}
        missing = required - headers.keys()
        if missing:
            raise ValueError(f'Faltan columnas requeridas: {", ".join(sorted(missing))}')

        creados = actualizados = omitidos = 0
        avisos = []
        with transaction.atomic():
            for numero_fila, row in enumerate(worksheet.iter_rows(min_row=12, values_only=True), start=12):
                if not any(value not in (None, '') for value in row):
                    continue

                def value(header):
                    return row[headers[header] - 1] if header in headers else None

                def fecha_celda(header):
                    return fecha(value(header), avisos, numero_fila, header)

                alumno_nombre = nombre_completo(value('APELLIDO1'), value('APELLIDO2'), value('NOMBRE'))
                dni = texto(value('DNI/NIF/NIE'))
                if not alumno_nombre:
                    omitidos += 1
                    continue

                nacimiento = fecha_celda('F_NTO')
                dni_normalizado = normalizar_dni(dni)
                if dni_normalizado:
                    alumnos = [
                        alumno for alumno in Alumno.objects.filter(autoescuela=autoescuela)
                        if alumno.dni_normalizado == dni_normalizado
                    ]
                else:
                    alumnos = list(Alumno.objects.filter(
                        autoescuela=autoescuela,
                        nombre__iexact=alumno_nombre,
                        fecha_nacimiento=nacimiento,
                    ))

                if len(alumnos) > 1:
                    raise ValueError(f'Hay varios alumnos para {alumno_nombre} ({dni or "sin DNI"}).')

                defaults = {
                    'nombre': alumno_nombre,
                    'nombre_pila': texto(value('NOMBRE')),
                    'apellido1': texto(value('APELLIDO1')),
                    'apellido2': texto(value('APELLIDO2')),
                    'dni': dni,
                    'permiso': texto(value('PMS')),
                    'numero_registro': numero_registro(value('N_REG')),
                    'fecha_alta': fecha_celda('F_ALTA'),
                    'fecha_nacimiento': nacimiento,
                    'fecha_inicio': fecha_celda('F.INI'),
                    'fecha_fin': fecha_celda('F.FIN'),
                    'causa': texto(value('CAUSA')),
                    'observaciones': texto(value('OBSERVACIONES')),
                    'fecha_apto_teorico': fecha_celda('F.APT.TEOR'),
                }
                if alumnos:
                    alumno = alumnos[0]
                    for field, field_value in defaults.items():
                        setattr(alumno, field, field_value)
                    alumno.save()
                    actualizados += 1
                else:
                    Alumno.objects.create(autoescuela=autoescuela, **defaults)
                    creados += 1

            if dry_run:
                transaction.set_rollback(True)

        return creados, actualizados, omitidos, avisos
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('archivo', type=Path, help='Ruta al LIBRO REGISTRO .xlsx')
    parser.add_argument('--autoescuela', required=True, help='Nombre exacto de la sede')
    parser.add_argument('--dry-run', action='store_true', help='Valida sin guardar cambios')
    args = parser.parse_args()

    creados, actualizados, omitidos, avisos = importar(args.archivo, args.autoescuela, args.dry_run)
    accion = 'Validación terminada' if args.dry_run else 'Importación terminada'
    print(f'{accion}: {creados} creados, {actualizados} actualizados, {omitidos} omitidos.')
    for aviso in avisos:
        print(f'AVISO: {aviso}')


if __name__ == '__main__':
    main()
